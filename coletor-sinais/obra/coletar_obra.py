#!/usr/bin/env python3
"""
Coletor de sinais precoces — classe 7 (obra física).

Porquê existe: um projeto industrial pede ligação à rede elétrica ANOS antes de
haver obra, encomendas ou notícias. É um compromisso administrativo datado, público
e assinado por uma entidade com nome — cabe na definição de sinal L2/L3 do motor
("permits obtidos") e é o sinal mais precoce que existe para qualquer transformação
que consuma energia.

Validação: o teste retroativo de 06/08/2026 mostrou que este sinal virou no
1.º semestre de 2024 (20,6 -> 58,0 GW de gás num semestre), quando o operador só
detetou o tema em outubro de 2025. ~18 meses de antecedência.

DUAS FONTES, dois papéis:

  1. EIA-860M (PRIMÁRIA, mensal) — geradores PLANEADOS, com entidade proprietária,
     estado, autoridade de balanço, tecnologia e data de operação prevista. Série
     mensal disponível desde 2023-01. É a fonte que define a latência do sistema.
     Inclui a folha "Canceled or Postponed": o sinal de REVERSÃO (falsificador).

  2. LBNL (SECUNDÁRIA, anual) — a fila de interconexão completa (38k pedidos desde
     2000). Publicação anual, logo latência de ~5 meses; serve para história
     profunda e para contexto, não para deteção.

O que se mede: capacidade de ENERGIA FIRME planeada por região e por entidade.
Firme = gás, nuclear, geotérmica, carvão — o que serve carga 24/7. Solar e eólica
não sinalizam procura industrial contínua da mesma forma.

O sinal NÃO é o nível — é a DERIVADA:
  - que região ganhou GW firmes planeados desde o retrato anterior;
  - que ENTIDADES aparecem pela primeira vez (difusão: a mesma lógica das 3-8
     empresas distintas do coletor de vocabulário, aplicada a betão e turbinas);
  - que capacidade passou a "cancelada/adiada" (reversão).

Uso:
  python3 coletar_obra.py                    # último mês + derivada vs retratos anteriores
  python3 coletar_obra.py --backfill         # constrói a série mensal toda (2023-01 -> hoje)
  python3 coletar_obra.py --mes 2025-06      # um mês específico
  python3 coletar_obra.py --lbnl             # análise de história profunda (fila LBNL)
"""
import argparse, json, os, re, sys, tempfile, urllib.request
from collections import defaultdict
from datetime import date
from pathlib import Path

BASE = Path(__file__).parent
SNAPS = BASE / 'snapshots'
XLSX_LBNL = BASE / 'lbnl_queue.xlsx'

MESES = ['january', 'february', 'march', 'april', 'may', 'june',
         'july', 'august', 'september', 'october', 'november', 'december']
# O mês corrente vive em /xls/; TODOS os anteriores em /archive/xls/ — por isso o
# arquivo tenta-se primeiro (acerta em 41 dos 42 casos e poupa uma redireção lenta).
EIA_URLS = ['https://www.eia.gov/electricity/data/eia860m/archive/xls/{m}_generator{y}.xlsx',
            'https://www.eia.gov/electricity/data/eia860m/xls/{m}_generator{y}.xlsx']
# As páginas do LBNL têm Cloudflare (403); os ficheiros em /sites/default/files/ não.
URL_LBNL = 'https://emp.lbl.gov/sites/default/files/2026-05/LBNL_Ix_Queue_Data_File_thru2025.xlsx'

LIMIAR_ALARME = 50.0     # % de crescimento a 12 meses da capacidade firme planeada
FIRME = re.compile(r'natural gas|nuclear|geothermal|coal|petroleum|other gas', re.I)
CARGA_BASE = re.compile(r'natural gas|nuclear', re.I)   # o subconjunto que responde a procura nova
UA = {'User-Agent': 'Mozilla/5.0'}


# ---------------------------------------------------------------- EIA-860M

class _SemRedirecao(urllib.request.HTTPRedirectHandler):
    """Um mês inexistente é redirecionado para a homepage do EIA com HTTP 200.
    Seguir a redireção custa 67 KB de HTML e vários segundos, e o resultado
    parece um sucesso. Tratamos qualquer redireção como 'não existe'."""
    def redirect_request(self, *a, **kw):
        return None


_OPENER = urllib.request.build_opener(_SemRedirecao)


def url_do_mes(ano, mes):
    """Devolve o primeiro URL que serve mesmo um xlsx (e não a homepage por redireção)."""
    for tpl in EIA_URLS:
        u = tpl.format(m=MESES[mes - 1], y=ano)
        try:
            r = _OPENER.open(urllib.request.Request(u, method='HEAD', headers=UA), timeout=20)
            if 'sheet' in r.headers.get('Content-Type', ''):
                return u
        except Exception:
            pass
    return None


def ler_folhas(caminho, folhas):
    """Lê folhas do xlsx em modo read-only (openpyxl direto: ~5x mais rápido que pandas)."""
    import openpyxl
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    saida = {}
    for nome in folhas:
        if nome not in wb.sheetnames:
            continue
        ws = wb[nome]
        linhas = ws.iter_rows(values_only=True)
        for _ in range(2):            # duas linhas de cabeçalho descritivo
            next(linhas, None)
        cols = [str(c).strip() if c is not None else '' for c in (next(linhas, None) or [])]
        idx = {c: i for i, c in enumerate(cols)}
        saida[nome] = (idx, [r for r in linhas if r and r[0] is not None])
    wb.close()
    return saida


def agregar(idx, linhas):
    """Extrai os agregados de uma folha (Planned ou Canceled or Postponed)."""
    def g(r, c):
        i = idx.get(c)
        return r[i] if i is not None and i < len(r) else None

    tot = {'unidades': 0, 'mw': 0.0}
    firme = {'unidades': 0, 'mw': 0.0}
    por_ba, por_ano, por_entidade, por_estado = defaultdict(float), defaultdict(float), defaultdict(float), defaultdict(float)
    ent_unidades = defaultdict(int)

    for r in linhas:
        try:
            mw = float(g(r, 'Nameplate Capacity (MW)') or 0)
        except (TypeError, ValueError):
            mw = 0.0
        tot['unidades'] += 1
        tot['mw'] += mw
        tec = str(g(r, 'Technology') or '')
        if not FIRME.search(tec):
            continue
        firme['unidades'] += 1
        firme['mw'] += mw
        ba = str(g(r, 'Balancing Authority Code') or '??').strip()
        ent = str(g(r, 'Entity Name') or '??').strip()
        est = str(g(r, 'Plant State') or '??').strip()
        por_ba[ba] += mw
        por_entidade[ent] += mw
        ent_unidades[ent] += 1
        por_estado[est] += mw
        ano = g(r, 'Planned Operation Year')
        if ano:
            try:
                por_ano[str(int(ano))] += mw
            except (TypeError, ValueError):
                pass

    top = sorted(por_entidade.items(), key=lambda kv: -kv[1])[:40]
    return {
        'total': {'unidades': tot['unidades'], 'gw': round(tot['mw'] / 1000, 2)},
        'firme': {'unidades': firme['unidades'], 'gw': round(firme['mw'] / 1000, 2)},
        'por_ba_gw': {k: round(v / 1000, 2) for k, v in sorted(por_ba.items(), key=lambda kv: -kv[1])[:20]},
        'por_estado_gw': {k: round(v / 1000, 2) for k, v in sorted(por_estado.items(), key=lambda kv: -kv[1])[:20]},
        'por_ano_gw': {k: round(v / 1000, 2) for k, v in sorted(por_ano.items())},
        'entidades_gw': {k: round(v / 1000, 3) for k, v in top},
        'entidades_unidades': {k: ent_unidades[k] for k, _ in top},
        'n_entidades_firmes': len(por_entidade),
    }


def colher_mes(ano, mes, forcar=False):
    """Descarrega, agrega e guarda o retrato de um mês. O xlsx é descartado (13MB)."""
    alvo = SNAPS / f'eia860m-{ano}-{mes:02d}.json'
    if alvo.exists() and not forcar:
        return json.loads(alvo.read_text())
    u = url_do_mes(ano, mes)
    if not u:
        return None
    with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as fh:
        tmp = fh.name
    try:
        req = urllib.request.Request(u, headers=UA)
        with urllib.request.urlopen(req, timeout=180) as resp, open(tmp, 'wb') as out:
            while chunk := resp.read(1 << 20):
                out.write(chunk)
        folhas = ler_folhas(tmp, ['Planned', 'Canceled or Postponed'])
    finally:
        os.unlink(tmp)

    snap = {'mes': f'{ano}-{mes:02d}', 'fonte': u, 'colhido_em': date.today().isoformat()}
    if 'Planned' in folhas:
        snap['planeado'] = agregar(*folhas['Planned'])
    if 'Canceled or Postponed' in folhas:
        snap['cancelado_adiado'] = agregar(*folhas['Canceled or Postponed'])
    SNAPS.mkdir(exist_ok=True)
    alvo.write_text(json.dumps(snap, ensure_ascii=False, indent=1))
    return snap


def carregar_serie():
    return [json.loads(p.read_text()) for p in sorted(SNAPS.glob('eia860m-*.json'))]


def relatar(serie):
    """A leitura que interessa: níveis, derivadas e entidades novas."""
    if not serie:
        print("Sem retratos. Correr com --backfill para construir a série.")
        return
    ult = serie[-1]
    p = ult.get('planeado', {})
    print(f"\n=== EIA-860M · {ult['mes']} ===")
    print(f"  Planeado total: {p.get('total', {}).get('unidades', 0):,} unidades · {p.get('total', {}).get('gw', 0):,.1f} GW")
    print(f"  Dos quais FIRME: {p.get('firme', {}).get('unidades', 0):,} · {p.get('firme', {}).get('gw', 0):,.1f} GW"
          f"  ({p.get('n_entidades_firmes', 0)} entidades distintas)")
    c = ult.get('cancelado_adiado', {}).get('firme', {})
    print(f"  Cancelado/adiado firme (reversão): {c.get('unidades', 0):,} · {c.get('gw', 0):,.1f} GW")

    print("\n--- Firme planeado por ano de operação (GW) ---")
    print('  ' + ' · '.join(f"{a}: {g:.1f}" for a, g in list(p.get('por_ano_gw', {}).items())[:8]))

    # Limiar de alarme, fixado a 06/08/2026 sobre a série 2023-01..2026-06.
    # Sobre essa série teria disparado em 2025-01, ~9 meses antes de o motor detetar
    # o tema da energia. RESSALVA: foi escolhido DEPOIS de ver a série — só disparos
    # futuros contam como validação.
    if len(serie) > 12:
        base = serie[-13].get('planeado', {}).get('firme', {}).get('gw', 0)
        agora = p.get('firme', {}).get('gw', 0)
        if base:
            var = (agora / base - 1) * 100
            estado = "** ALARME **" if var >= LIMIAR_ALARME else "abaixo do limiar"
            print(f"\n[{estado}] variação a 12 meses: {var:+.0f}% (limiar: +{LIMIAR_ALARME:.0f}%)")

    for jan, rot in ((1, '1 mês'), (3, '3 meses'), (12, '12 meses')):
        if len(serie) <= jan:
            continue
        ant = serie[-1 - jan]
        pa = ant.get('planeado', {})
        d = p.get('firme', {}).get('gw', 0) - pa.get('firme', {}).get('gw', 0)
        base = pa.get('firme', {}).get('gw', 0)
        pct = f" ({d / base * 100:+.0f}%)" if base else ""
        print(f"\n--- Derivada a {rot} (vs {ant['mes']}): {d:+.1f} GW firmes{pct} ---")
        ba_now, ba_ant = p.get('por_ba_gw', {}), pa.get('por_ba_gw', {})
        deltas = sorted(((b, ba_now.get(b, 0) - ba_ant.get(b, 0)) for b in set(ba_now) | set(ba_ant)),
                        key=lambda kv: -abs(kv[1]))[:6]
        for b, dd in deltas:
            antes = ba_ant.get(b, 0)
            flag = "  ** DUPLICOU" if antes > 1 and ba_now.get(b, 0) > 2 * antes else ""
            print(f"    {b:<6} {antes:6.1f} -> {ba_now.get(b, 0):6.1f} GW  ({dd:+.1f}){flag}")

    if len(serie) > 6:
        ant = serie[-7]
        novas = [e for e in p.get('entidades_gw', {}) if e not in ant.get('planeado', {}).get('entidades_gw', {})]
        if novas:
            print(f"\n--- Entidades no top-40 que não lá estavam há 6 meses ({ant['mes']}) ---")
            for e in novas[:12]:
                print(f"    {e}  ·  {p['entidades_gw'][e]:.2f} GW  ·  {p['entidades_unidades'][e]} unidades")

    print("\n--- Top entidades por GW firmes planeados (semente de watchlist) ---")
    for e, g in list(p.get('entidades_gw', {}).items())[:12]:
        print(f"    {e:<52} {g:6.2f} GW · {p['entidades_unidades'][e]} un.")
    print("\nLeitura: o sinal é a DERIVADA e a difusão de entidades, não o nível.")
    print("Um sinal aqui é um POINTER para a fase de descoberta — nunca uma tese.")


# ---------------------------------------------------------------- LBNL

def relatar_lbnl(desde, descarregar):
    import pandas as pd, warnings
    warnings.filterwarnings('ignore')
    if descarregar or not XLSX_LBNL.exists():
        print(f"A descarregar de {URL_LBNL[:60]}...")
        XLSX_LBNL.write_bytes(urllib.request.urlopen(
            urllib.request.Request(URL_LBNL, headers=UA), timeout=180).read())
        print(f"  {XLSX_LBNL.stat().st_size:,} bytes")
    df = pd.read_excel(XLSX_LBNL, sheet_name='03. Complete Queue Data', skiprows=1)
    df.columns = [str(c).strip() for c in df.columns]
    df['mw'] = pd.to_numeric(df['mw_1'], errors='coerce').fillna(0)
    df['q_date'] = pd.to_datetime(df['q_date'], errors='coerce')
    f = df[df.type_clean.astype(str).str.contains('Gas|Nuclear|Geothermal|Coal', case=False, na=False)].dropna(subset=['q_date'])
    f = f[f.q_date.dt.year >= desde]
    f['sem'] = f.q_date.dt.year.astype(str) + '-S' + ((f.q_date.dt.month > 6).astype(int) + 1).astype(str)
    print(f"\n=== LBNL (história profunda) — energia firme pedida por semestre desde {desde} ===")
    tot = f.groupby('sem').agg(pedidos=('mw', 'size'), gw=('mw', lambda x: x.sum() / 1000))
    prev = None
    for s, r in tot.iterrows():
        d = f"  ({r.gw / prev:+.1f}x)" if prev and prev > 0 else ""
        alerta = "  ** DUPLICACAO" if prev and r.gw > 2 * prev else ""
        print(f"  {s}: {int(r.pedidos):4} pedidos · {r.gw:6.1f} GW{d}{alerta}")
        prev = r.gw


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--backfill', action='store_true', help='constrói a série mensal desde 2023-01')
    ap.add_argument('--mes', help='colher um mês específico (AAAA-MM)')
    ap.add_argument('--forcar', action='store_true', help='recolher mesmo que o retrato exista')
    ap.add_argument('--lbnl', action='store_true', help='análise de história profunda (fila LBNL)')
    ap.add_argument('--desde', type=int, default=2021, help='ano inicial da análise LBNL')
    args = ap.parse_args()

    if args.lbnl:
        relatar_lbnl(args.desde, descarregar=not XLSX_LBNL.exists())
        return

    SNAPS.mkdir(exist_ok=True)
    if args.mes:
        a, m = map(int, args.mes.split('-'))
        if colher_mes(a, m, args.forcar):
            print(f"Retrato {args.mes} guardado.")
        else:
            print(f"{args.mes} indisponível no EIA.", file=sys.stderr)
        relatar(carregar_serie())
        return

    hoje = date.today()
    if args.backfill:
        import concurrent.futures as cf
        alvos = [(a, m) for a in range(2023, hoje.year + 1) for m in range(1, 13)
                 if (a, m) <= (hoje.year, hoje.month)
                 and not ((SNAPS / f'eia860m-{a}-{m:02d}.json').exists() and not args.forcar)]
        print(f"{len(alvos)} meses em falta. A colher (6 em paralelo)...", flush=True)
        with cf.ThreadPoolExecutor(6) as ex:      # 6: rápido sem martelar o servidor do EIA
            fut = {ex.submit(colher_mes, a, m, args.forcar): (a, m) for a, m in alvos}
            for f in cf.as_completed(fut):
                a, m = fut[f]
                try:
                    r = f.result()
                except Exception as e:
                    r, e_ = None, e
                    print(f"  {a}-{m:02d}: ERRO {e_}", flush=True)
                    continue
                print(f"  {a}-{m:02d}: " + (f"{r['planeado']['firme']['gw']:.1f} GW firmes planeados"
                                            if r else "indisponível"), flush=True)
    else:
        # Só se procuram meses POSTERIORES ao retrato mais recente que já existe:
        # o EIA publica uma vez por mês, logo a maioria das corridas não tem nada
        # novo para colher e não deve gastar tempo a sondar o servidor.
        serie = sorted(SNAPS.glob('eia860m-*.json'))
        a, m = (2023, 1)
        if serie:
            a, m = map(int, serie[-1].stem.split('-')[1:])
            m += 1
            if m == 13:
                a, m = a + 1, 1
        novos = 0
        while (a, m) <= (hoje.year, hoje.month):
            if colher_mes(a, m):
                print(f"  novo retrato: {a}-{m:02d}")
                novos += 1
            m += 1
            if m == 13:
                a, m = a + 1, 1
        if not novos:
            print("Sem edição nova do EIA desde o último retrato — leitura sobre a série existente.")
    relatar(carregar_serie())


if __name__ == '__main__':
    main()
